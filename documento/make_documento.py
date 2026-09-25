import os

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

RAIZ = r"D:\Parcial1-IoT"
VERDE = RGBColor(0x2F, 0x6F, 0x4E)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
for lado in ("left_margin", "right_margin"):
    setattr(sec, lado, Cm(2.2))
sec.top_margin = sec.bottom_margin = Cm(2.2)

estilo = doc.styles["Normal"]
estilo.font.name = "Calibri"
estilo.font.size = Pt(11)
for nombre, tam in (("Heading 1", 16), ("Heading 2", 13)):
    h = doc.styles[nombre]
    h.font.name = "Calibri"
    h.font.size = Pt(tam)
    h.font.color.rgb = VERDE


def sombrear(celda, color):
    tc = celda._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), color)
    tc.append(sh)


def tabla(encabezados, filas, anchos=None, tam=9):
    t = doc.add_table(rows=1, cols=len(encabezados))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, e in enumerate(encabezados):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(e)
        r.bold = True
        r.font.size = Pt(tam)
        r.font.color.rgb = RGBColor(255, 255, 255)
        sombrear(c, "2F6F4E")
    for fila in filas:
        celdas = t.add_row().cells
        for i, v in enumerate(fila):
            celdas[i].text = ""
            celdas[i].paragraphs[0].add_run(str(v)).font.size = Pt(tam)
    if anchos:
        for fila in t.rows:
            for i, a in enumerate(anchos):
                fila.cells[i].width = Cm(a)
    doc.add_paragraph()
    return t


def parrafo(texto, negrita=False):
    p = doc.add_paragraph()
    r = p.add_run(texto)
    r.bold = negrita
    return p


def figura(ruta, ancho=16.5, pie=None):
    doc.add_picture(ruta, width=Cm(ancho))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if pie:
        p = doc.add_paragraph(pie)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].italic = True
        p.runs[0].font.size = Pt(9)


p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Cm(2)
doc.add_picture(os.path.join(RAIZ, "recursos", "logo_cacao_pixel.png"), width=Cm(5))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

for texto, tam, neg in (
    ("Universidad Autónoma de Bucaramanga", 14, False),
    ("Internet de las Cosas", 14, False),
    ("", 10, False),
    ("Parcial 1", 26, True),
    ("Escenario IoT Central con flota heterogénea de 10 dispositivos", 15, False),
    ("Granja y cultivo de cacao - San Vicente de Chucuri", 15, True),
    ("", 10, False),
    ("Rafael Antonio Arias Monsalve", 13, False),
    ("Docente: Javier Pinzón", 12, False),
    ("Septiembre de 2026", 12, False),
):
    q = doc.add_paragraph()
    q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = q.add_run(texto)
    r.font.size = Pt(tam)
    r.bold = neg
    if neg:
        r.font.color.rgb = VERDE

doc.add_page_break()

doc.add_heading("Historial de versiones", 1)
tabla(
    ["Versión", "Fecha", "Descripción"],
    [
        ["0.1", "21/09/2026", "Elección del escenario 5.3 y catálogo de 10 dispositivos y 10 orígenes"],
        ["0.2", "22/09/2026", "Aplicación de IoT Central, 8 plantillas DTDL y primer dispositivo en línea (simulador nativo)"],
        ["0.3", "23/09/2026", "Siete orígenes automáticos en una VM, panel Cuarto de control y dos reglas de alerta"],
        ["0.4", "24/09/2026", "Primer documento de avance con telemetría de dos días"],
    ],
    [2.5, 3, 11.5],
    10,
)

doc.add_heading("1. Escenario", 1)
parrafo(
    "El escenario elegido es el 5.3, Granja y cultivo de cacao. La finca está en la zona de San Vicente de Chucurí "
    "(latitud 6.88, longitud -73.41). Al ser un predio rural no hay red institucional: la salida a Internet se hace "
    "por un módem 4G/LTE instalado en la caseta central, y los nodos de campo se conectan por Wi-Fi local a ese módem."
)
parrafo(
    "Las variables indispensables son la humedad de suelo por lote, la temperatura y humedad del dosel de sombra, la "
    "meteorología del predio, el nivel del tanque de riego y la temperatura de la caja de fermentación. Además hay "
    "un nodo de calidad de aire y un nodo de perímetro en la bodega de secado."
)

doc.add_heading("2. Arquitectura de referencia", 1)
parrafo(
    "Se tiene la siguiente arquitectura para la infraestructura digital de la finca de cacao, teniendo en cuenta los "
    "dispositivos, los métodos de comunicación y los servicios de la nube. En esta versión del proyecto no existe una "
    "infraestructura física, por lo que los sensores son simulados, aunque la arquitectura describe su implementación "
    "física."
)
doc.add_heading("Arquitectura de referencia", 2)
figura(os.path.join(RAIZ, "documento", "diagrama_arquitectura.png"), 17.2, "Figura 1. Arquitectura de referencia de la finca de cacao.")
parrafo(
    "La arquitectura incluye los dispositivos de la finca, que se comunican por Wi-Fi inalámbrico o por Ethernet según "
    "su ubicación; los datos pasan por un gateway (el módem 4G/LTE de la caseta) y salen a Internet. Los dos nodos que "
    "consumen APIs públicas salen directo a Internet, sin pasar por el módem de campo."
)
parrafo(
    "Mediante la conexión a Internet los datos llegan a Azure IoT Central, donde el servicio de aprovisionamiento (DPS) "
    "asigna cada dispositivo al IoT Hub. Desde ahí los datos siguen la ruta caliente, tibia y fría descrita en la "
    "arquitectura del servicio, sobre una capa de servicios PaaS que aporta disponibilidad, escalabilidad y "
    "recuperación ante desastres. En la experiencia web de administración se gestionan los dispositivos (datos sin "
    "procesar, estado de conectividad, modelado y trabajos), se visualizan y analizan los datos (paneles, analítica y "
    "reglas) y se administran los usuarios y las organizaciones."
)
parrafo(
    "En la integración empresarial, Azure Maps (Atlas Weather) se consulta para la calidad del aire del nodo 07 y "
    "las reglas de alerta envían correos electrónicos."
)

doc.add_heading("Sensores y dispositivos", 2)
parrafo("Teniendo en cuenta el diseño de la finca, se tienen 10 dispositivos:")
for item in (
    "3 lotes de cultivo con sensores de humedad y temperatura de suelo, conductividad e iluminancia.",
    "1 nodo de dosel de sombra con temperatura, humedad relativa e iluminancia.",
    "1 estación de campo con pluviómetro y humedad foliar.",
    "1 nodo de meteorología del predio y 1 nodo de calidad de aire, alimentados por APIs.",
    "1 caja de secado y fermentación con temperatura, humedad y masa.",
    "1 reservorio de riego con nivel de tanque, caudal y bomba.",
    "1 nodo de perímetro y bodega con puerta, movimiento y temperatura.",
):
    doc.add_paragraph(item, style="List Bullet")
parrafo("En la sección de plantillas se describe el resumen del DTDL de los dispositivos con sus variables.")

doc.add_heading("Comunicación", 2)
parrafo("Cada grupo de dispositivos tiene un tipo de conexión distinto:")
for item in (
    "Los nodos de campo (lotes, dosel, estación y reservorio) se conectan por Wi-Fi al módem 4G/LTE de la caseta, "
    "porque están dispersos en el predio y no se puede tender cable hasta cada uno.",
    "La caja de secado y el nodo de la bodega van por Ethernet, porque están fijos dentro de la caseta y la bodega.",
    "Los nodos de meteorología y calidad de aire no son sensores del predio: consultan APIs públicas y llegan por "
    "Internet directo, desde un equipo con salida propia.",
    "Cada dispositivo se registra con DPS usando una clave derivada de la clave de grupo de inscripción (HMAC-SHA256 "
    "del identificador). Los orígenes usan protocolos distintos: MQTT sobre TLS (8883), MQTT sobre WebSockets (443) "
    "y HTTPS.",
):
    doc.add_paragraph(item, style="List Bullet")

doc.add_heading("3. Catálogo de dispositivos y orígenes", 1)
tabla(
    ["#", "Dispositivo", "Plantilla", "Origen de envío", "Variables enviadas"],
    [
        ["01", "lote-cultivo-01", "NodoCultivo", "Digital Twin / simulador nativo de IoT Central", "humedadSuelo, temperaturaSuelo, conductividad, iluminanciaPAR"],
        ["02", "lote-cultivo-02", "NodoCultivo", "ESP32 en Wokwi (DHT22, potenciómetros, LDR)", "humedadSuelo, temperaturaSuelo, conductividad, iluminanciaPAR"],
        ["03", "lote-cultivo-03", "NodoCultivo", "Python con azure-iot-device (MQTT)", "humedadSuelo, iluminanciaPAR"],
        ["04", "dosel-sombra-01", "NodoDoselSombra", "Python con paho-mqtt, SAS manual", "temperatura, humedadRelativa, iluminancia"],
        ["05", "estacion-campo-01", "NodoEstacionCampo", "Node.js con azure-iot-device (MQTT)", "lluviaLote, humedadFoliar"],
        ["06", "meteorologia-predio-01", "NodoMeteorologia", "API pública Open-Meteo", "temperatura, humedadRelativa, lluvia, viento, radiacion"],
        ["07", "calidad-aire-01", "NodoCalidadAire", "Atlas Weather (Azure Maps Weather API)", "pm25, humedadRelativa, aqi"],
        ["08", "secado-fermentacion-01", "NodoProceso", "Python con azure-iot-device sobre WebSockets", "temperaturaCaja, humedadRelativa, masaEstimada"],
        ["09", "reservorio-riego-01", "NodoRiego", "ESP32 en Wokwi (ultrasónico y relé)", "nivelTanque, caudal"],
        ["10", "perimetro-bodega-01", "NodoPerimetro", "REST manual: HTTPS directo al IoT Hub", "temperaturaBodega, puerta, movimiento"],
    ],
    [1, 3.4, 3, 5, 5],
    8.5,
)

doc.add_heading("4. Parámetros y hojas de datos de referencia", 1)
parrafo(
    "Los dispositivos de Wokwi y los nodos simulados representan sensores reales. La tabla indica el sensor de "
    "referencia de cada variable, su rango y la hoja de datos consultada."
)
tabla(
    ["Variable", "Sensor de referencia", "Rango / precisión", "Código en el proyecto"],
    [
        ["humedadSuelo", "Sensor capacitivo de humedad de suelo v2.0", "0-100 % tras calibrar", "sketch.ino (ADC GPIO34), lote03_python_sdk.py"],
        ["temperaturaSuelo / temperatura", "DHT22 (AM2302)", "-40 a 80 °C, ±0,5 °C", "sketch.ino (GPIO15), dosel04_mqtt.py"],
        ["conductividad", "Sensor de conductividad de suelo", "0-4 mS/cm en el modelo", "sketch.ino (ADC GPIO35)"],
        ["iluminancia / iluminanciaPAR", "BH1750 / fotoresistor LDR", "1-65535 lux", "sketch.ino (ADC GPIO32), dosel04_mqtt.py"],
        ["lluviaLote", "Pluviómetro de cubeta basculante", "0,2 mm por pulso", "estacion05_node.js"],
        ["humedadFoliar", "Sensor resistivo de humedad foliar", "0-100 %", "estacion05_node.js"],
        ["nivelTanque", "Ultrasónico HC-SR04 (simulado) / JSN-SR04T", "2-400 cm / 20-600 cm", "reservorio-riego-01/sketch.ino"],
        ["caudal", "Sensor de flujo YF-S201", "1-30 L/min", "reservorio-riego-01/sketch.ino"],
        ["temperaturaCaja / temperaturaBodega", "DS18B20", "-55 a 125 °C, ±0,5 °C", "secado08_websockets.py, perimetro10_rest.py"],
        ["masaEstimada", "Celda de carga con HX711", "0-5 kg", "secado08_websockets.py"],
        ["puerta, movimiento", "Contacto magnético; PIR HC-SR501", "Digital 0/1", "perimetro10_rest.py"],
        ["pm25, aqi", "Azure Maps Weather (calidad de aire)", "Índice global", "aire07_atlas.py"],
        ["temperatura, lluvia, viento, radiacion", "Open-Meteo API", "Datos horarios de modelo", "meteo06_openmeteo.py"],
    ],
    [3.6, 4.6, 3.6, 5.6],
    8.5,
)

doc.add_heading("5. Plantillas de dispositivo (DTDL)", 1)
parrafo("Se publicaron 8 plantillas, una por rol de nodo. Los archivos están en la carpeta modelo del repositorio.")
tabla(
    ["Plantilla", "Telemetría", "Propiedades", "Comandos"],
    [
        ["NodoCultivo", "humedadSuelo, temperaturaSuelo, conductividad, iluminanciaPAR", "intervaloMuestreoSeg (escribible, 300), loteId", "-"],
        ["NodoDoselSombra", "temperatura, humedadRelativa, iluminancia", "intervaloMuestreoSeg (120)", "-"],
        ["NodoEstacionCampo", "lluviaLote, humedadFoliar", "intervaloMuestreoSeg (600)", "-"],
        ["NodoMeteorologia", "temperatura, humedadRelativa, lluvia, viento, radiacion", "fuenteDatos", "-"],
        ["NodoCalidadAire", "pm25, humedadRelativa, aqi", "fuenteDatos", "-"],
        ["NodoProceso", "temperaturaCaja, humedadRelativa, masaEstimada", "loteEnProceso (escribible)", "-"],
        ["NodoRiego", "nivelTanque, caudal", "nivelTanqueMin (escribible, 20), bombaActiva", "activarBomba(duracionMin)"],
        ["NodoPerimetro", "temperaturaBodega, puerta, movimiento", "-", "-"],
    ],
    [3.3, 6, 5, 3.2],
    8.5,
)

doc.add_heading("6. Asincronía, desconexiones y operación en vivo", 1)
parrafo(
    "Cada script simula apagones con un ciclo propio para que los dispositivos no se desconecten a la vez. En los "
    "servicios de la VM el dispositivo cierra la conexión, deja pasar la ventana de apagón y se reconecta solo. "
    "Los intervalos de envío son distintos (45 s, 60 s, 90 s, 120 s, 300 s, 600 s y 900 s)."
)
tabla(
    ["Dispositivo", "Intervalo", "Ciclo de apagón (cada / dura)"],
    [
        ["lote-cultivo-03", "60 s", "240 min / 25 min"],
        ["dosel-sombra-01", "45 s", "300 min / 40 min"],
        ["estacion-campo-01", "300 s", "260 min / 35 min"],
        ["secado-fermentacion-01", "120 s", "200 min / 30 min"],
        ["perimetro-bodega-01", "90 s", "180 min / 35 min"],
        ["meteorologia-predio-01", "600 s", "sin apagón programado"],
        ["calidad-aire-01", "900 s", "sin apagón programado"],
    ],
    [6, 3, 7],
    9.5,
)

doc.add_heading("7. Cuarto de control", 1)
parrafo(
    "El panel Cuarto de control - Finca Cacao muestra el conteo de la flota, siete indicadores (promedio, máximo y "
    "mínimo de humedad de suelo, nivel mínimo del tanque, temperatura máxima de la caja, PM2.5 máximo y lluvia "
    "acumulada) y doce gráficas de línea con las variables de los ocho tipos de dispositivo. Se creó con un script "
    "que usa la API de IoT Central (panel_api.py)."
)
doc.add_heading("Reglas de alerta", 2)
tabla(
    ["Regla", "Plantilla", "Condición", "Acción"],
    [
        ["Fermentacion caja caliente", "NodoProceso", "temperaturaCaja > 42", "Correo"],
        ["Bodega temperatura alta", "NodoPerimetro", "temperaturaBodega > 28", "Correo"],
    ],
    [5, 3.5, 4.5, 3],
    9.5,
)

doc.add_heading("8. Telemetría por día", 1)
resumen = pd.read_csv(os.path.join(RAIZ, "datos", "resumen_por_dia.csv"))
dias = sorted(resumen["dia"].unique())
piv = resumen.pivot_table(index="dispositivo", columns="dia", values="horas_con_datos", aggfunc="sum").fillna(0).astype(int)
parrafo(
    "La tabla cuenta las horas del día en las que llegó al menos un mensaje, con hora de Colombia (UTC-5). "
    "Los datos salen de la consulta a IoT Central a través de su API."
)
tabla(["Dispositivo"] + [d[8:10] + "/" + d[5:7] for d in dias], [[i] + [piv.loc[i, d] if d in piv.columns else 0 for d in dias] for i in piv.index], None, 9.5)
parrafo("Los días 24 y 25 de septiembre se completan cuando termine la captura.")

graficas = [
    ("01_humedad_suelo.png", "Humedad del suelo por lote"),
    ("02_temperatura_suelo.png", "Temperatura del suelo"),
    ("03_dosel_temperatura.png", "Temperatura bajo sombra"),
    ("04_meteo_temperatura.png", "Temperatura de la API Open-Meteo"),
    ("05_aire_pm25.png", "PM2.5 desde Atlas Weather"),
    ("06_fermentacion_temperatura.png", "Temperatura de la caja de fermentación"),
    ("07_bodega_temperatura.png", "Temperatura de la bodega"),
    ("08_estacion_humedad_foliar.png", "Humedad foliar de la estación de campo"),
]
doc.add_heading("Gráficas de la telemetría recibida", 2)
for i, (archivo, pie) in enumerate(graficas, 1):
    ruta = os.path.join(RAIZ, "evidencias", "graficas", archivo)
    if os.path.exists(ruta):
        figura(ruta, 16.5, f"Figura {i + 1}. {pie}.")

doc.add_heading("9. Repositorio", 1)
parrafo(
    "El repositorio contiene el README, los modelos DTDL (modelo), los scripts de cada origen (scripts), los proyectos de "
    "Wokwi (wokwi), los datos exportados (datos) y las evidencias. Las claves y el ID scope no se suben: se leen de "
    "variables de entorno y de archivos secrets.h locales, que están en el .gitignore."
)

doc.add_heading("10. Sustentación", 1)
parrafo(
    "Se ejecutan dos códigos en dos equipos distintos: un script de Python en el portátil y el ESP32 de Wokwi en el "
    "navegador, mostrando en IoT Central los estados de conexión y la telemetría en vivo."
)

prop = doc.core_properties
prop.author = "Rafael Antonio Arias Monsalve"
prop.last_modified_by = "Rafael Antonio Arias Monsalve"
prop.comments = ""
prop.title = "Parcial 1 - Internet de las Cosas"
doc.save(os.path.join(RAIZ, "documento", "Parcial1_IoT_Central_v3.docx"))
print("listo")
